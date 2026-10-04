"""Render the Kazakh deck and editable PowerPoint from its HTML source.

Requires Python Playwright, python-pptx, Pillow, PyMuPDF and Chromium.
Use --capture to refresh real prototype screenshots before rendering.
"""
import argparse
from functools import partial
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from threading import Thread
import json

import fitz
from PIL import Image, ImageOps, ImageDraw
from playwright.sync_api import sync_playwright, expect
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
parser = argparse.ArgumentParser()
parser.add_argument('--capture', action='store_true')
args = parser.parse_args()

class QuietHandler(SimpleHTTPRequestHandler):
    def log_message(self, *args):
        pass

def rgb(value):
    import re
    parts = re.findall(r'\d+', value)
    return RGBColor(*map(int, parts[:3]))

server = ThreadingHTTPServer(('127.0.0.1', 0), partial(QuietHandler, directory=str(ROOT)))
Thread(target=server.serve_forever, daemon=True).start()
base = f'http://127.0.0.1:{server.server_port}'
assets = ROOT / 'docs/presentation-assets'
assets.mkdir(exist_ok=True)
try:
    with sync_playwright() as p:
        browser = p.chromium.launch(executable_path='/usr/bin/chromium', args=['--no-sandbox', '--disable-dev-shm-usage'])
        page = browser.new_page(viewport={'width':1440, 'height':900})
        if args.capture:
            page.goto(base + '/site/')
            expect(page.get_by_role('heading', name='Обзор производства', exact=True)).to_be_visible()
            page.screenshot(path=str(assets / 'overview.png'))
            page.locator('.content-grid').screenshot(path=str(assets / 'process-risk.png'))
            page.get_by_role('button', name='Отчёты', exact=True).click()
            page.get_by_role('slider').fill('30')
            page.locator('#hourly-downtime-cost').fill('100000')
            page.locator('#intervention-cost').fill('20000')
            expect(page.locator('.cost-metrics')).to_be_visible()
            assert '105000' in ''.join(page.locator('.cost-metrics strong').all_text_contents()).replace('\u00a0', '').replace(' ', '')
            page.locator('.cost-evaluation').screenshot(path=str(assets / 'economics-30.png'))
            page.locator('.chart').screenshot(path=str(assets / 'comparison-chart.png'))
            page.get_by_role('button', name='Настройки', exact=True).click()
            page.get_by_role('heading', name='Источник данных', exact=True).locator('..').locator('..').locator('..').screenshot(path=str(assets / 'csv-import.png'))
        page.set_viewport_size({'width':1280, 'height':720})
        page.goto(base + '/docs/presentation.html', wait_until='networkidle')
        page.evaluate('document.fonts.ready')
        assert page.locator('.slide').count() == 10
        assert page.evaluate('Array.from(document.images).every(i => i.complete && i.naturalWidth > 0)')
        layout = page.evaluate('''() => Array.from(document.querySelectorAll('.slide')).map(slide => {
            const base = slide.getBoundingClientRect();
            const rect = el => {const b=el.getBoundingClientRect(); return {x:b.x-base.x,y:b.y-base.y,w:b.width,h:b.height};};
            return {title:slide.dataset.label, bg:getComputedStyle(slide).backgroundColor,
                boxes:Array.from(slide.querySelectorAll('.box,.formula,.tag')).map(el=>({...rect(el),bg:getComputedStyle(el).backgroundColor})),
                images:Array.from(slide.querySelectorAll('img')).map(el=>({...rect(el),src:el.getAttribute('src')})),
                texts:Array.from(slide.querySelectorAll('h1,h2,h3,p,.eyebrow,.number,.formula,.tag,.flow>.box,.arrow,.checklist li,.footer span,a.link')).map(el=>{
                    const s=getComputedStyle(el);return {...rect(el),text:el.innerText,color:s.color,font:parseFloat(s.fontSize),line:parseFloat(s.lineHeight)||parseFloat(s.fontSize)*1.2,bold:parseInt(s.fontWeight)>=600,footer:!!el.closest('.footer')};
                })};
        })''')
        overflow = [(i+1,t['text'][:60]) for i,s in enumerate(layout) for t in s['texts'] if not t['footer'] and (t['y']+t['h']>664 or t['x']+t['w']>1221)]
        assert not overflow, overflow
        for i in [0,2,6]:
            page.locator('.slide').nth(i).screenshot(path=str(assets / f'slide-{i+1:02d}.png'))
        page.pdf(path=str(ROOT / 'docs/presentation.pdf'), prefer_css_page_size=True, print_background=True)
        browser.close()
finally:
    server.shutdown()
    server.server_close()

# Editable text and screenshot elements, matching HTML geometry at 96 px/inch.
prs = Presentation()
prs.slide_width = Inches(1280/96)
prs.slide_height = Inches(720/96)
for data in layout:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(data['bg'])
    for box in data['boxes']:
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, *[Inches(box[k]/96) for k in ['x','y','w','h']])
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(box['bg'])
        shape.line.fill.background()
    for im in data['images']:
        path = ROOT / 'docs' / im['src']
        with Image.open(path) as source:
            scale = min(im['w']/source.width, im['h']/source.height)
            w,h = source.width*scale,source.height*scale
        slide.shapes.add_picture(str(path), Inches((im['x']+(im['w']-w)/2)/96), Inches((im['y']+(im['h']-h)/2)/96), Inches(w/96), Inches(h/96))
    for text in data['texts']:
        shape = slide.shapes.add_textbox(Inches(text['x']/96), Inches(text['y']/96), Inches(text['w']/96), Inches((text['h']+6)/96))
        frame = shape.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        for i,line in enumerate(text['text'].split('\n')):
            para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            para.text = line
            para.font.name = 'DejaVu Sans'
            para.font.size = Pt(text['font']*.75)
            para.font.bold = text['bold']
            para.font.color.rgb = rgb(text['color'])
            para.space_before = para.space_after = Pt(0)
            para.line_spacing = Pt(text['line']*.75)
prs.save(ROOT / 'docs/presentation-kk.pptx')

doc = fitz.open(ROOT / 'docs/presentation.pdf')
assert len(doc) == 10
words = '\n'.join(page.get_text() for page in doc)
for name in ['Zhubanov team', 'Бижан Әлижан', 'Қанымбек Марат', 'Құлтас Баубек', 'Жұбанов университеті', 'Ақтөбе қаласы']:
    assert name in words, name
assert 'Перед сдачей добавить' not in words and '\ufffd' not in words
thumbs = []
for i, page in enumerate(doc):
    pix = page.get_pixmap(matrix=fitz.Matrix(.4,.4))
    image = Image.frombytes('RGB',[pix.width,pix.height],pix.samples)
    thumbs.append(ImageOps.pad(image,(512,288),color='#dde3e8'))
sheet = Image.new('RGB',(1052,1530),'#dde3e8')
draw = ImageDraw.Draw(sheet)
for i,thumb in enumerate(thumbs):
    x,y = 8+(i%2)*524,8+(i//2)*304
    sheet.paste(thumb,(x,y))
    draw.text((x,y+289),str(i+1),fill='#192330')
sheet.save(assets / 'contact-sheet.png')
report = {'language':'kk','slides':10,'teamSize':3,'realScreenshots':5,'pdfTextAndNamesPassed':True,'layoutBoundsPassed':True,'editablePowerPoint':True,'interfaceLanguage':'ru','applicationTestsRerun':False}
(ROOT / 'docs/verification/presentation-kk.json').write_text(json.dumps(report,indent=2,ensure_ascii=False)+'\n')
print(json.dumps(report,ensure_ascii=False))
